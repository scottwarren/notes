from pathlib import Path
from docx import Document
from docx.shared import Inches, Pt
from docx.enum.text import WD_ALIGN_PARAGRAPH

base=Path('/Users/scottwarren/projects/notes/Personal/Visa/801 - Permanent Residency')
md=base/'Prepared Documents/Form 80 Review Items - Kristina Aleksandrova Warren - WaKr-2631-801-3.md'
out=base/'Prepared Documents/Form 80 Review Items - Kristina Aleksandrova Warren - WaKr-2631-801-3.docx'
d=Document()
sec=d.sections[0]; sec.top_margin=Inches(.65); sec.bottom_margin=Inches(.65); sec.left_margin=Inches(.72); sec.right_margin=Inches(.72)
styles=d.styles
styles['Normal'].font.name='Arial'; styles['Normal'].font.size=Pt(10)
styles['Title'].font.name='Arial'; styles['Title'].font.size=Pt(18)
styles['Heading 1'].font.name='Arial'; styles['Heading 1'].font.size=Pt(13)
for line in md.read_text().splitlines():
    if line.startswith('# '):
        d.add_heading(line[2:],0)
    elif line.startswith('## '):
        d.add_heading(line[3:],1)
    elif line.startswith('- [ ] '):
        p=d.add_paragraph(style='List Bullet'); p.add_run('☐ '+line[6:])
    elif line.startswith('- [x] '):
        p=d.add_paragraph(style='List Bullet'); p.add_run('☑ '+line[6:])
    elif line.strip():
        d.add_paragraph(line)
    else:
        d.add_paragraph('')
head=sec.header.paragraphs[0]; head.text='Kristina Aleksandrova Warren — Form 80 review — WaKr-2631-801-3'; head.alignment=WD_ALIGN_PARAGRAPH.CENTER
for p in sec.footer.paragraphs:
    p.alignment=WD_ALIGN_PARAGRAPH.CENTER
    p.text='Review draft — 18 October 2026'
d.core_properties.title='Form 80 Review Items - Kristina Aleksandrova Warren'
d.save(out)
print(out)
