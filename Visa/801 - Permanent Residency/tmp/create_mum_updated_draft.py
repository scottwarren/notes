from pathlib import Path
from docx import Document
from docx.shared import Pt
from docx.oxml.ns import qn

src=Path('Prepared Documents/Evidence Gathered/Rosita Paragreen Witness Supporting Statement - Original Draft - 2026-10-17 - WaKr-2631-801-3.docx')
out=Path('Prepared Documents/Rosita Paragreen Witness Supporting Statement - Updated Draft - WaKr-2631-801-3.docx')
doc=Document(src)
t=doc.tables[0]

t.cell(0,1).text=("I, Rosita Paragreen, of 9 Woollahra Court, Varsity Lakes, Queensland 4227, "
                  "Retired, formerly a lawyer,\nmake the following declaration under the Statutory Declarations Act 1959:")

blocks=[
("Statement of Declaration:", True),
("I make the following statement in support of Kristina Aleksandrova Warren's permanent Partner visa application.", False),
("Visa Applicant's Details:\nFamily Name: Warren\nGiven Names: Kristina Aleksandrova\nDuration of Acquaintance: Approximately 9 years, since early 2017", True),
("Applicant's Spouse:\nFamily Name: Warren\nGiven Names: Scott Phillip\nDuration of Acquaintance: Since his birth in 1991, approximately 35 years\nNature of Relationship: Scott Phillip Warren is my son and Kristina Aleksandrova Warren is my daughter-in-law.", True),
("How I know Scott and Kristina:", True),
("1. I first became acquainted with Kristina in early 2017 through WhatsApp messages and video calls. I met her in person in early 2019 when she came to Australia for a family wedding, and we celebrated Scott and Kristina's engagement with our extended family.", False),
("2. Kristina and Scott stayed with me during visits to Australia, including Christmas 2022. I visited them in their home in Sofia in 2022 and again in 2023 for their wedding. I met Kristina's family on both visits and remain in contact with her mother through WhatsApp or Messenger.", False),
("3. I spent Christmas 2024 with Scott and Kristina at their home in Urunga, near Coffs Harbour. Kristina, Scott, Kristina's mother and her mother's partner stayed with me at my home for a short holiday in April 2025. I also travelled with Scott and Kristina to Bulgaria in December 2025 and spent Christmas with Kristina's family.", False),
("4. I remain in regular contact with Scott and Kristina by telephone, video calls and messages. Through our ongoing contact and family visits, I continue to observe them supporting each other and planning their future together as a married couple.", False),
("Belief in the Relationship:", True),
("5. I believe Kristina and Scott's marital relationship is genuine and continuing. They have been together since at least 2017, travelled extensively, established a home together in Sofia, and married in Bulgaria in July 2023 after dating and living together for several years.", False),
("6. I have observed their commitment during their visits to Australia and my visits to Bulgaria. Both families and extended families have embraced them as a couple, and their shared life and regular communication demonstrate a committed marriage.", False),
("7. They moved to New South Wales in late 2024 and now live together in Coffs Harbour. I continue to see them make decisions together about their home, work and future.", False),
("Additional Information:", True),
("8. As Scott's mother, I fully support Kristina and Scott's permanent Partner visa application. My relationship with them has allowed me to observe their relationship over many years, both in person and through regular contact.", False),
("Contact Details:\nContact Number: 0405 470 508\nEmail Address: rparagreen18@gmail.com", True),
("Further Instructions: Attach one of the following identity documents, certified by the same eligible person witnessing this statutory declaration:\n• Driver's licence\n• Passport biographical page\n• Australian Permanent Resident Visa Grant Letter (no certification required)\n• Australian Citizenship Certificate", False),
("I understand that a person who intentionally makes a false statement in a statutory declaration is guilty of an offence under section 11 of the Statutory Declarations Act 1959, and I believe that the statements in this declaration are true in every particular.", False),
]
cell=t.cell(1,1)
cell._element.clear_content()
for text,bold in blocks:
    p=cell.add_paragraph(text)
    p.paragraph_format.space_after=Pt(4)
    for run in p.runs:
        run.font.name='Arial'; run._element.rPr.rFonts.set(qn('w:eastAsia'),'Arial'); run.font.size=Pt(9); run.font.bold=bold

# Remove template placeholder numbers and leave usable signature fields.
t.cell(2,1).text='Signature: ______________________________________________\n\n'
t.cell(3,1).text='Declared at ______________________________ on ______ of __________________ 20____\n\nBefore me,\n'
t.cell(4,1).text='Signature of witness: ___________________________________\n\n'
t.cell(5,1).text='Full name: ______________________________________________\nQualification: ____________________________________________\nAddress: _________________________________________________\n_________________________________________________________\n'

for table in doc.tables:
    for row in table.rows:
        for c in row.cells:
            for p in c.paragraphs:
                for run in p.runs:
                    run.font.name='Arial'; run._element.rPr.rFonts.set(qn('w:eastAsia'),'Arial')
                    if run.font.size is None: run.font.size=Pt(9)

doc.core_properties.title='Rosita Paragreen Witness Supporting Statement - Updated Draft - WaKr-2631-801-3'
doc.core_properties.subject='Subclass 801 partner visa witness statement for review'
doc.save(out)
print(out)
