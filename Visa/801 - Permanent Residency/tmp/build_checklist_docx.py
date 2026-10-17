from pathlib import Path
import re

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK, WD_COLOR_INDEX
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


SOURCE = Path("/Users/scottwarren/projects/notes/Personal/Visa/801 - Permanent Residency/Visa 801 - Document Checklist.md")
OUTPUT = Path("/Users/scottwarren/projects/notes/Personal/Visa/801 - Permanent Residency/Prepared Documents/Visa 801 Document Checklist.docx")


def set_font(run, name: str, size: float | None = None) -> None:
    run.font.name = name
    run._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), name)
    run._element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), name)
    if size is not None:
        run.font.size = Pt(size)


def set_cell_or_paragraph_shading(paragraph, fill: str) -> None:
    p_pr = paragraph._p.get_or_add_pPr()
    shading = p_pr.find(qn("w:shd"))
    if shading is None:
        shading = OxmlElement("w:shd")
        p_pr.append(shading)
    shading.set(qn("w:fill"), fill)


def add_page_number(paragraph) -> None:
    paragraph.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    prefix = paragraph.add_run("Page ")
    set_font(prefix, "Aptos", 9)
    run = paragraph.add_run()
    set_font(run, "Aptos", 9)
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = " PAGE "
    separate = OxmlElement("w:fldChar")
    separate.set(qn("w:fldCharType"), "separate")
    result = OxmlElement("w:t")
    result.text = "1"
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    run._r.extend([begin, instr, separate, result, end])


def add_inline_markup(paragraph, text: str, base_size: float = 10.5) -> None:
    token_re = re.compile(r"(\*\*.*?\*\*|`.*?`)")
    for part in token_re.split(text):
        if not part:
            continue
        if part.startswith("**") and part.endswith("**"):
            run = paragraph.add_run(part[2:-2])
            run.bold = True
            set_font(run, "Aptos", base_size)
        elif part.startswith("`") and part.endswith("`"):
            run = paragraph.add_run(part[1:-1])
            set_font(run, "Aptos Mono", base_size - 0.5)
        else:
            run = paragraph.add_run(part)
            set_font(run, "Aptos", base_size)


def configure_styles(doc: Document) -> None:
    normal = doc.styles["Normal"]
    normal.font.name = "Aptos"
    normal._element.rPr.rFonts.set(qn("w:ascii"), "Aptos")
    normal._element.rPr.rFonts.set(qn("w:hAnsi"), "Aptos")
    normal.font.size = Pt(10.5)
    normal.paragraph_format.space_after = Pt(5)
    normal.paragraph_format.line_spacing = 1.08

    title = doc.styles["Title"]
    title.font.name = "Aptos Display"
    title._element.rPr.rFonts.set(qn("w:ascii"), "Aptos Display")
    title._element.rPr.rFonts.set(qn("w:hAnsi"), "Aptos Display")
    title.font.size = Pt(23)
    title.font.bold = True
    title.font.color.rgb = RGBColor(0, 0, 0)
    title.paragraph_format.space_after = Pt(12)

    for style_name, size, before, after in [
        ("Heading 1", 15, 14, 6),
        ("Heading 2", 12, 10, 4),
        ("Heading 3", 11, 8, 3),
    ]:
        style = doc.styles[style_name]
        style.font.name = "Aptos Display"
        style._element.rPr.rFonts.set(qn("w:ascii"), "Aptos Display")
        style._element.rPr.rFonts.set(qn("w:hAnsi"), "Aptos Display")
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = RGBColor(0, 0, 0)
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)
        style.paragraph_format.keep_with_next = True


def build() -> None:
    doc = Document()
    configure_styles(doc)

    section = doc.sections[0]
    section.top_margin = Inches(0.65)
    section.bottom_margin = Inches(0.6)
    section.left_margin = Inches(0.72)
    section.right_margin = Inches(0.72)

    header = section.header
    hp = header.paragraphs[0]
    hp.text = "Visa 801 document status    Client reference WaKr-2631-801-3"
    hp.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    for run in hp.runs:
        set_font(run, "Aptos", 8.5)
        run.font.color.rgb = RGBColor(90, 90, 90)

    footer = section.footer
    footer.paragraphs[0].text = ""

    lines = SOURCE.read_text(encoding="utf-8").splitlines()
    current_section = ""
    for raw in lines:
        if not raw.strip():
            continue

        if raw.startswith("# "):
            p = doc.add_paragraph(style="Title")
            add_inline_markup(p, raw[2:].strip(), 23)
            for run in p.runs:
                run.bold = True
                run.font.color.rgb = RGBColor(0, 0, 0)
            continue

        if raw.startswith("## "):
            current_section = raw[3:].strip()
            p = doc.add_paragraph(style="Heading 1")
            add_inline_markup(p, current_section, 15)
            continue

        if raw.startswith("### "):
            p = doc.add_paragraph(style="Heading 2")
            add_inline_markup(p, raw[4:].strip(), 12)
            continue

        stripped = raw.lstrip("\t ")
        indent_units = len(raw) - len(stripped)
        if stripped.startswith("* ["):
            match = re.match(r"\* \[([ xX])\] (.*)", stripped)
            if not match:
                continue
            checked = match.group(1).lower() == "x"
            body = match.group(2)
            depth = 1 if indent_units else 0
            p = doc.add_paragraph()
            p.paragraph_format.left_indent = Inches(0.22 + depth * 0.27)
            p.paragraph_format.first_line_indent = Inches(-0.18)
            p.paragraph_format.space_after = Pt(3.5)
            p.paragraph_format.keep_together = False

            status = p.add_run("DONE  " if checked else "TO DO  ")
            status.bold = True
            set_font(status, "Aptos", 9.5)
            if checked:
                status.font.color.rgb = RGBColor(26, 112, 61)
            else:
                status.font.color.rgb = RGBColor(156, 72, 0)
                if "[P1]" in body or current_section == "Remaining priority actions":
                    status.font.highlight_color = WD_COLOR_INDEX.YELLOW
            add_inline_markup(p, body, 10.2)
            if body.startswith("[P1] Joint ownership or joint rental"):
                p.paragraph_format.keep_with_next = True
            if body.startswith("Lease agreement"):
                p.paragraph_format.keep_with_next = True
            continue

        if stripped.startswith("* "):
            body = stripped[2:]
            depth = 1 if indent_units else 0
            p = doc.add_paragraph()
            p.paragraph_format.left_indent = Inches(0.22 + depth * 0.27)
            p.paragraph_format.first_line_indent = Inches(-0.16)
            bullet = p.add_run("•  ")
            set_font(bullet, "Aptos", 10.5)
            add_inline_markup(p, body, 10.2)
            continue

        p = doc.add_paragraph()
        if raw.startswith("Resolved —") or raw.startswith("Resolved -"):
            lead, sep, rest = raw.partition("—") if "—" in raw else raw.partition("-")
            run = p.add_run(lead + (" — " if sep else " "))
            run.bold = True
            run.font.color.rgb = RGBColor(26, 112, 61)
            set_font(run, "Aptos", 10.2)
            add_inline_markup(p, rest.strip(), 10.2)
        else:
            add_inline_markup(p, raw, 10.2)
        if raw.startswith("Updated 18 October 2026"):
            p.paragraph_format.space_after = Pt(8)
        if raw.startswith("801 folder:") or raw.startswith("Evidence:"):
            for run in p.runs:
                run.font.size = Pt(9)

    doc.core_properties.title = "Visa 801 Document Checklist"
    doc.core_properties.subject = "Current completion status and remaining actions"
    doc.core_properties.keywords = "Visa 801, checklist, WaKr-2631-801-3"
    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    doc.save(OUTPUT)
    print(OUTPUT)


if __name__ == "__main__":
    build()
