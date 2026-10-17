from __future__ import annotations

import re
from pathlib import Path

from PIL import Image
from docx import Document
from docx.enum.section import WD_ORIENT
from docx.enum.table import WD_ALIGN_VERTICAL, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path("/Users/scottwarren/projects/notes/Personal/Visa/801 - Permanent Residency")
PREPARED = ROOT / "Prepared Documents"
PHOTO_DIR = Path("/Users/scottwarren/Documents/Visas/Permanent Resident Visa/801 - Permanent Visa/Evidence/Relationship Photo Timeline - WaKr-2631-801-3")
REL_TEMPLATE = PREPARED / "Scott Relationship Statement 801 Initial Draft.docx"
REL_MD = PREPARED / "Scott Relationship Statement 801 First Review Draft.md"
REL_OUT = PREPARED / "Scott Relationship Statement 801 First Review Draft - WaKr-2631-801-3.docx"
MUM_MD = PREPARED / "Mum Witness Statement Review Notes - WaKr-2631-801-3.md"
MUM_OUT = PREPARED / "Mum Witness Statement Review Notes - WaKr-2631-801-3.docx"
PHOTO_OUT = PREPARED / "Relationship Photo Timeline - August 2025 to July 2026 - WaKr-2631-801-3.docx"


PHOTO_CAPTIONS = [
    ("15 August 2025", "Buccan, Queensland", "Friends' wedding ceremony attended by Scott and Kristina together."),
    ("15 August 2025", "Buccan, Queensland", "Scott and Kristina together at their friends' wedding."),
    ("29 August 2025", "North Wall Beach, New South Wales", "Scott and Kristina on a walk together at the beach."),
    ("8 November 2025", "Calga, New South Wales", "Scott and Kristina together during a day trip."),
    ("8 November 2025", "Calga, New South Wales", "Scott and Kristina watching kangaroos at sunset during their day trip."),
    ("23 December 2025", "Borovets, Bulgaria", "Scott and Kristina celebrating Christmas with both sides of their family."),
    ("23 December 2025", "Borovets, Bulgaria", "Relaxing together in the chalet hot tub during the family Christmas trip."),
    ("23 December 2025", "Borovets, Bulgaria", "Scott and Kristina with family on the ski lift."),
    ("23 December 2025", "Borovets, Bulgaria", "Scott and Kristina with family on the ski lift during their Christmas holiday."),
    ("23 December 2025", "Borovets, Bulgaria", "Scott, Kristina and family together in the snow."),
    ("25 December 2025", "Borovets, Bulgaria", "Christmas Day meal with Scott, Kristina and family from both sides."),
    ("28 December 2025", "Sofia, Bulgaria", "Scott and Kristina sharing a family dinner during their Christmas visit."),
    ("4 January 2026", "Ovoshtnik, Bulgaria", "Scott and Kristina together during their holiday in Bulgaria."),
    ("28 February 2026", "Engie Stadium, Sydney", "Scott and Kristina attending a concert together."),
    ("16 June 2026", "Sofia, Bulgaria", "Scott and Kristina with family before returning to Australia."),
    ("16 June 2026", "Sofia, Bulgaria", "Scott and Kristina spending time with family during the June visit."),
    ("11 July 2026", "Woolgoolga, New South Wales", "Scott and Kristina on a coastal walk together."),
    ("11 July 2026", "Woolgoolga, New South Wales", "Scott and Kristina together on the Woolgoolga coast."),
]


def set_font(run, name="Aptos", size=None, bold=None, italic=None, color=None):
    run.font.name = name
    run._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), name)
    run._element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), name)
    if size is not None:
        run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold
    if italic is not None:
        run.italic = italic
    if color is not None:
        run.font.color.rgb = RGBColor(*color)


def remove_cell_content(cell):
    tc = cell._tc
    for child in list(tc):
        if child.tag != qn("w:tcPr"):
            tc.remove(child)
    cell.add_paragraph()


def set_repeat_table_header(row):
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    tr_pr.append(tbl_header)


def set_cell_margins(cell, top=90, start=110, bottom=90, end=110):
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for tag, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{tag}"))
        if node is None:
            node = OxmlElement(f"w:{tag}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def build_relationship_statement():
    doc = Document(REL_TEMPLATE)
    cell = doc.tables[0].cell(1, 1)
    remove_cell_content(cell)
    lines = REL_MD.read_text(encoding="utf-8").splitlines()
    in_statement = False
    for raw in lines:
        text = raw.strip()
        if text == "## History of our relationship and home":
            in_statement = True
        if text == "## Items Scott still needs to confirm before signing":
            break
        if not in_statement or not text:
            continue
        if text.startswith("## "):
            p = cell.add_paragraph()
            p.paragraph_format.space_before = Pt(6)
            p.paragraph_format.space_after = Pt(2)
            p.paragraph_format.keep_with_next = True
            r = p.add_run(text[3:])
            set_font(r, "Arial", 9.5, bold=True)
        elif re.match(r"^\d+\.\s", text):
            p = cell.add_paragraph()
            p.paragraph_format.space_after = Pt(3)
            p.paragraph_format.line_spacing = 1.0
            r = p.add_run(text)
            set_font(r, "Arial", 8.8)

    p = cell.add_paragraph()
    p.paragraph_format.space_before = Pt(5)
    p.paragraph_format.space_after = Pt(0)
    r = p.add_run(
        "I understand that a person who intentionally makes a false statement in a statutory declaration "
        "is guilty of an offence under section 11 of the Statutory Declarations Act 1959, and I believe "
        "that the statements in this declaration are true in every particular."
    )
    set_font(r, "Arial", 8.8)

    first_cell = doc.tables[0].cell(0, 1)
    for para in first_cell.paragraphs:
        if "Scott Phillip Warren" in para.text:
            para.text = (
                "I, 1 Scott Phillip Warren, 8/4 Prince Street, Coffs Harbour, NSW 2450, Web Developer,\n\n"
                "make the following declaration under the Statutory Declarations Act 1959:"
            )
            for run in para.runs:
                set_font(run, "Arial", 9)

    doc.core_properties.title = "Scott Relationship Statement 801 First Review Draft"
    doc.core_properties.subject = "Partner visa relationship statement for review"
    doc.core_properties.keywords = "Visa 801, relationship statement, WaKr-2631-801-3"
    doc.save(REL_OUT)


def configure_simple_doc(doc, title_text, subtitle):
    section = doc.sections[0]
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(0.72)
    section.bottom_margin = Inches(0.68)
    section.left_margin = Inches(0.8)
    section.right_margin = Inches(0.8)

    normal = doc.styles["Normal"]
    normal.font.name = "Aptos"
    normal._element.rPr.rFonts.set(qn("w:ascii"), "Aptos")
    normal._element.rPr.rFonts.set(qn("w:hAnsi"), "Aptos")
    normal.font.size = Pt(10.5)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.08

    title = doc.styles["Title"]
    title.font.name = "Aptos Display"
    title._element.rPr.rFonts.set(qn("w:ascii"), "Aptos Display")
    title._element.rPr.rFonts.set(qn("w:hAnsi"), "Aptos Display")
    title.font.size = Pt(22)
    title.font.bold = True
    title.font.color.rgb = RGBColor(0, 0, 0)
    title.paragraph_format.space_after = Pt(6)

    for name, size in (("Heading 1", 15), ("Heading 2", 12)):
        style = doc.styles[name]
        style.font.name = "Aptos Display"
        style._element.rPr.rFonts.set(qn("w:ascii"), "Aptos Display")
        style._element.rPr.rFonts.set(qn("w:hAnsi"), "Aptos Display")
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = RGBColor(0, 0, 0)
        style.paragraph_format.space_before = Pt(10)
        style.paragraph_format.space_after = Pt(4)
        style.paragraph_format.keep_with_next = True

    p = doc.add_paragraph(style="Title")
    p.add_run(title_text)
    p = doc.add_paragraph()
    r = p.add_run(subtitle)
    set_font(r, size=9.5, color=(85, 85, 85))
    p.paragraph_format.space_after = Pt(12)


def add_inline_bold(paragraph, text):
    for part in re.split(r"(\*\*.*?\*\*)", text):
        if not part:
            continue
        bold = part.startswith("**") and part.endswith("**")
        content = part[2:-2] if bold else part
        r = paragraph.add_run(content)
        set_font(r, size=10.5, bold=bold)


def build_mum_review():
    doc = Document()
    configure_simple_doc(
        doc,
        "Mum Witness Statement Review Notes",
        "Review date 18 October 2026    Client reference WaKr-2631-801-3",
    )
    lines = MUM_MD.read_text(encoding="utf-8").splitlines()
    for raw in lines[1:]:
        text = raw.strip()
        if not text or text.startswith("Review date:") or text.startswith("Client reference:") or text.startswith("Source reviewed:"):
            continue
        if text.startswith("## "):
            doc.add_paragraph(text[3:], style="Heading 1")
        elif re.match(r"^\d+\.\s", text):
            p = doc.add_paragraph(style="List Number")
            add_inline_bold(p, re.sub(r"^\d+\.\s", "", text))
        elif text.startswith("“"):
            p = doc.add_paragraph()
            p.paragraph_format.left_indent = Inches(0.25)
            p.paragraph_format.right_indent = Inches(0.25)
            r = p.add_run(text)
            set_font(r, size=10.2, italic=True)
        else:
            p = doc.add_paragraph()
            add_inline_bold(p, text)
    doc.core_properties.title = "Mum Witness Statement Review Notes"
    doc.core_properties.subject = "Review of Rosita Paragreen witness statement draft"
    doc.core_properties.keywords = "Visa 801, witness statement, WaKr-2631-801-3"
    doc.save(MUM_OUT)


def natural_photo_key(path):
    return int(path.stem.rsplit("-", 1)[-1].strip())


def photo_size(path, max_w, max_h):
    with Image.open(path) as image:
        w, h = image.size
    scale = min(max_w / w, max_h / h)
    return Inches(w * scale / 96), Inches(h * scale / 96)


def build_photo_timeline():
    doc = Document()
    configure_simple_doc(
        doc,
        "Relationship Photo Timeline August 2025 to July 2026",
        "Scott Phillip Warren and Kristina Aleksandrova Warren    Client reference WaKr-2631-801-3",
    )
    p = doc.add_paragraph(
        "These photographs show family events, travel and ordinary time together. The August 2025 photographs provide brief context immediately before the Subclass 820 grant on 24 September 2025; the remaining photographs cover the period after grant."
    )
    p.paragraph_format.space_after = Pt(10)

    paths = sorted(PHOTO_DIR.glob("*.jpeg"), key=natural_photo_key)
    if len(paths) != len(PHOTO_CAPTIONS):
        raise RuntimeError(f"Expected {len(PHOTO_CAPTIONS)} photos, found {len(paths)}")

    for idx, (path, (date, location, caption)) in enumerate(zip(paths, PHOTO_CAPTIONS)):
        if idx and idx % 2 == 0:
            doc.add_page_break()
        table = doc.add_table(rows=2, cols=1)
        table.alignment = WD_TABLE_ALIGNMENT.CENTER
        table.autofit = False
        for row in table.rows:
            row.cells[0].width = Inches(6.75)
            row.cells[0].vertical_alignment = WD_ALIGN_VERTICAL.CENTER
            set_cell_margins(row.cells[0], top=65, start=65, bottom=65, end=65)

        image_p = table.cell(0, 0).paragraphs[0]
        image_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        width, height = photo_size(path, 6.45 * 96, 3.45 * 96)
        image_p.add_run().add_picture(str(path), width=width, height=height)

        caption_p = table.cell(1, 0).paragraphs[0]
        caption_p.alignment = WD_ALIGN_PARAGRAPH.LEFT
        r = caption_p.add_run(f"{date}  {location}\n")
        set_font(r, size=9.5, bold=True)
        r = caption_p.add_run(caption)
        set_font(r, size=9.5)
        caption_p.paragraph_format.space_after = Pt(0)
        spacer = doc.add_paragraph()
        spacer.paragraph_format.space_after = Pt(1)

    doc.core_properties.title = "Relationship Photo Timeline August 2025 to July 2026"
    doc.core_properties.subject = "Relationship evidence for visa 801"
    doc.core_properties.keywords = "Visa 801, photo timeline, WaKr-2631-801-3"
    doc.save(PHOTO_OUT)


if __name__ == "__main__":
    PREPARED.mkdir(parents=True, exist_ok=True)
    build_relationship_statement()
    build_mum_review()
    build_photo_timeline()
    print(REL_OUT)
    print(MUM_OUT)
    print(PHOTO_OUT)
