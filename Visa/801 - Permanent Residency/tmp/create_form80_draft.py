from copy import deepcopy
from pathlib import Path

from docx import Document
from docx.enum.text import WD_COLOR_INDEX
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor
from pypdf import PdfReader, PdfWriter


ROOT = Path("/Users/scottwarren/projects/notes/Personal/Visa/801 - Permanent Residency")
SRC_PDF = ROOT / "Prepared Documents/Form 80 - Kristina Aleksandrova Warren - Working Source - WaKr-2631-801-3.pdf"
SRC_DOCX = ROOT / "Prepared Documents/Form 80 Attachment - Kristina Aleksandrova Warren - Working Source - WaKr-2631-801-3.docx"
OUT_PDF = ROOT / "Prepared Documents/Form 80 - Kristina Aleksandrova Warren - Review Draft - WaKr-2631-801-3.pdf"
OUT_DOCX = ROOT / "Prepared Documents/Form 80 Attachment - Kristina Aleksandrova Warren - Review Draft - WaKr-2631-801-3.docx"
NOTES_MD = ROOT / "Prepared Documents/Form 80 Review Items - Kristina Aleksandrova Warren - WaKr-2631-801-3.md"
NOTES_DOCX = ROOT / "Prepared Documents/Form 80 Review Items - Kristina Aleksandrova Warren - WaKr-2631-801-3.docx"


def clear_shading(cell):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is not None:
        tc_pr.remove(shd)


def shade(cell, fill="FFF2CC"):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell(cell, text, review=False):
    cell.text = text
    clear_shading(cell)
    if review:
        shade(cell)
    for p in cell.paragraphs:
        for r in p.runs:
            r.font.name = "Aptos"
            r.font.size = Pt(8.5)
            if review:
                r.font.highlight_color = WD_COLOR_INDEX.YELLOW


def set_row(row, values, review_cols=()):
    for i, value in enumerate(values):
        set_cell(row.cells[i], value, i in review_cols)


def insert_after(table, row_index):
    source = table.rows[row_index]._tr
    new_tr = deepcopy(source)
    source.addnext(new_tr)
    return table.rows[row_index + 1]


def update_attachment():
    doc = Document(SRC_DOCX)

    # Q17 address history: most recent first, with one date boundary still requiring confirmation.
    t = doc.tables[0]
    set_row(t.rows[1], ["28.07.2025", "Current", "8/4 Prince Street, Coffs Harbour, NSW 2450", "Australia"])
    set_row(t.rows[2], ["11.11.2024", "27.07.2025", "28 Fitzroy Street, Urunga, NSW 2455", "Australia"])
    set_row(t.rows[3], ["[REVIEW: exact move date]", "10.11.2024", "12 Millet Place, Upper Coomera, QLD 4209", "Australia"], review_cols=(0,))
    mum = insert_after(t, 3)
    set_row(mum, ["21.06.2024", "[REVIEW: exact move date]", "9 Woollahra Court, Varsity Lakes, QLD 4227", "Australia"], review_cols=(1,))

    # Q18 international travel: add trips after the last recorded 2024 trip.
    t = doc.tables[1]
    set_row(t.rows[1], ["03.06.2026", "18.06.2026", "Family bereavement and family visit", "Sofia / Bulgaria"])
    r = insert_after(t, 1)
    set_row(r, ["11.12.2025", "08.01.2026", "Family visit and holiday", "Sofia / Bulgaria"])
    r = insert_after(t, 2)
    set_row(r, ["30.09.2025", "10.10.2025", "Holiday", "Bali / Indonesia"])

    # Q19 employment/unemployment. Known periods are added; uncertain employment details are highlighted.
    t = doc.tables[2]
    set_row(
        t.rows[1],
        ["01.10.2026", "Current", "Coffs Coast Wildlife Sanctuary", "Wildlife sanctuary", "[REVIEW: exact role title and main duties]", "[REVIEW: employer street address]", "Australia"],
        review_cols=(4, 5),
    )
    set_row(
        t.rows[2],
        ["21.06.2024", "30.09.2026", "", "", "Not employed; supported by husband. Studied at TAFE NSW and volunteered with wildlife organisations. [REVIEW wording and any employment/volunteering dates]", "", "Australia"],
        review_cols=(4,),
    )
    r = insert_after(t, 2)
    set_row(r, ["03.10.2022", "20.06.2024", "", "", "Unemployed; supported by husband, travelled and pursued personal interests", "", "Bulgaria"])

    # Q20 Australian qualifications. Exact course dates/title details need Kristina's records.
    t = doc.tables[3]
    r = insert_after(t, 0)
    set_row(
        r,
        ["[REVIEW: start date]", "[REVIEW: exact completion date in June 2026]", "TAFE NSW", "Certificate III in Wildlife and Exhibited Animal Care [REVIEW official title/code]", "Completed", "Coffs Harbour Education Campus [REVIEW exact address]", "Australia"],
        review_cols=(0, 1, 3, 5),
    )
    r = insert_after(t, 1)
    set_row(
        r,
        ["[REVIEW: start date]", "[REVIEW: exact completion date in June 2025]", "TAFE NSW", "Certificate II in Animal Care (ACM20121-01V03) [REVIEW official title]", "Completed", "Coffs Harbour Education Campus [REVIEW exact address]", "Australia"],
        review_cols=(0, 1, 3, 5),
    )

    doc.core_properties.title = "Kristina Warren Form 80 Attachment - Review Draft"
    doc.core_properties.subject = "Updated for subclass 801 review; yellow items require confirmation"
    doc.save(OUT_DOCX)


def update_pdf():
    reader = PdfReader(SRC_PDF)
    writer = PdfWriter()
    writer.clone_document_from_reader(reader)
    updates = {
        "ap_fs_cities": "Coffs Harbour, NSW",
        "syncA_country_of_residence_p2": "DECEASED",
    }
    for page in writer.pages:
        try:
            writer.update_page_form_field_values(page, updates, auto_regenerate=False)
        except Exception:
            pass
    writer.set_need_appearances_writer(True)
    writer.add_metadata({
        "/Title": "Form 80 - Kristina Aleksandrova Warren - Review Draft",
        "/Subject": "Review draft for subclass 801; signature and date intentionally blank",
    })
    with OUT_PDF.open("wb") as f:
        writer.write(f)


REVIEW_SECTIONS = [
    ("Main Form 80 — confirm before signing", [
        "Confirm Kristina's mobile number remains +61 480 655 043.",
        "Confirm the primary email remains visa@scottandkristinawarren.com and list any other personal, student or work email address required by Question 15.",
        "Confirm the current Bulgarian passport details remain correct: passport 389306673, issued 11 September 2023 and expiring 11 September 2028.",
        "Confirm the Bulgarian identity-card number remains 652800497 and the birth-certificate number remains 0040.",
        "Confirm every Question 36 character answer remains 'No'. Any change must be disclosed accurately.",
        "Confirm the two prior visa-refusal explanations at Question 39 remain complete and accurate.",
        "Kristina's father has been updated from Bulgaria to DECEASED in the parent-details section. Confirm all other partner, parent, step-parent and sibling details remain current.",
        "Confirm Scott's and Rosita's contact details in Australia remain correct, including phone numbers and addresses.",
        "The further-stay city has been updated from Brisbane/Gold Coast to Coffs Harbour, NSW. Confirm this is the preferred answer.",
        "Do not sign or date Part S until Kristina and Affinity have reviewed the complete form and attachment.",
    ]),
    ("Question 17 — address history", [
        "Provide the exact date Kristina and Scott moved from 9 Woollahra Court, Varsity Lakes to 12 Millet Place, Upper Coomera. The draft leaves this boundary highlighted because Form 80 requires no gaps.",
        "Confirm 12 Millet Place was occupied until 10 November 2024, immediately before the Urunga tenancy began on 11 November 2024.",
        "Confirm the address sequence and dates: Varsity Lakes from 21 June 2024; Upper Coomera; Urunga from 11 November 2024 to 27 July 2025; Coffs Harbour from 28 July 2025 to current.",
    ]),
    ("Question 18 — international travel", [
        "Confirm the destination wording and dates for Bali: 30 September to 10 October 2025.",
        "Confirm the Bulgaria Christmas trip should be recorded as 11 December 2025 to 8 January 2026, using the Australian return-arrival date.",
        "Confirm the bereavement/family trip should be recorded as 3 June to 18 June 2026, using the Australian return-arrival date.",
        "Confirm there were no other international trips after 11 February 2024.",
    ]),
    ("Question 19 — employment and unemployment", [
        "Confirm Kristina's exact Coffs Coast Wildlife Sanctuary role title, employer street address and main duties.",
        "Confirm whether 1 October 2026 is the employment commencement date for Form 80 purposes, rather than only the first shift date.",
        "Review the 21 June 2024 to 30 September 2026 entry describing Kristina as not employed, supported by Scott, studying at TAFE NSW and volunteering. Add exact paid work or material volunteer periods if required.",
        "Provide short, accurate duty descriptions for Wildlife BG, Eyes On Four Paws Foundation, the waitress roles and Gadol Ltd. The agency added '[main duties]' prompts to the prior entries.",
    ]),
    ("Question 20 — education", [
        "Confirm exact commencement and completion dates for the Certificate II and Certificate III.",
        "Confirm the official Certificate III course title and code. The working description is Certificate III in Wildlife and Exhibited Animal Care.",
        "Confirm the official Certificate II title. Email evidence supports Certificate II in Animal Care, course ACM20121-01V03.",
        "Confirm the TAFE NSW campus name and street address used for both courses.",
    ]),
    ("Question 35 — visa history", [
        "Review the agency-prepared Partner (subclass 820) row and confirm whether Affinity wants the 24 September 2025 grant date recorded anywhere in the visa-history attachment.",
        "Confirm no other Australian visa or entry needs to be added after the two historical subclass 651 entries.",
    ]),
]


def create_review_notes():
    lines = [
        "# Form 80 Review Items — Kristina Aleksandrova Warren",
        "",
        "Prepared 18 October 2026. Yellow cells in the editable attachment need confirmation. The PDF retains the agency's previously completed answers, with only two known updates applied: Coffs Harbour as the further-stay city and Kristina's father marked DECEASED. This is a review draft and must not be signed yet.",
        "",
    ]
    for heading, items in REVIEW_SECTIONS:
        lines += [f"## {heading}", ""] + [f"- [ ] {x}" for x in items] + [""]
    NOTES_MD.write_text("\n".join(lines), encoding="utf-8")

    doc = Document()
    sec = doc.sections[0]
    sec.top_margin = Inches(0.65); sec.bottom_margin = Inches(0.65)
    sec.left_margin = Inches(0.75); sec.right_margin = Inches(0.75)
    normal = doc.styles["Normal"]
    normal.font.name = "Aptos"; normal.font.size = Pt(10.5)
    title = doc.add_heading("Form 80 Review Items", 0)
    title.add_run("\nKristina Aleksandrova Warren").font.size = Pt(15)
    p = doc.add_paragraph("Prepared 18 October 2026 · Client reference WaKr-2631-801-3")
    p.runs[0].font.color.rgb = RGBColor(90, 90, 90)
    p = doc.add_paragraph()
    r = p.add_run("REVIEW DRAFT — DO NOT SIGN YET")
    r.bold = True; r.font.highlight_color = WD_COLOR_INDEX.YELLOW
    doc.add_paragraph("Yellow cells in the editable attachment need confirmation. The PDF retains the agency's previously completed answers, with only two known updates applied: Coffs Harbour as the further-stay city and Kristina's father marked DECEASED.")
    for heading, items in REVIEW_SECTIONS:
        doc.add_heading(heading, level=1)
        for item in items:
            p = doc.add_paragraph(style="List Bullet")
            p.add_run("☐ ").bold = True
            p.add_run(item)
    doc.core_properties.title = "Form 80 Review Items - Kristina Aleksandrova Warren"
    doc.save(NOTES_DOCX)


if __name__ == "__main__":
    update_attachment()
    update_pdf()
    create_review_notes()
    print(OUT_PDF)
    print(OUT_DOCX)
    print(NOTES_DOCX)
