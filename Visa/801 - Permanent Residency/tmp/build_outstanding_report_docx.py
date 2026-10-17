from pathlib import Path
import re

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_COLOR_INDEX
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
SOURCE = ROOT / "Prepared Documents" / "Visa 801 Outstanding Items Report.md"
OUTPUT = ROOT / "Prepared Documents" / "Visa 801 Outstanding Items Report - WaKr-2631-801-3.docx"


def font(run, size=10.5, bold=None):
    run.font.name = "Aptos"
    run._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), "Aptos")
    run._element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), "Aptos")
    run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold


def inline(paragraph, text, size=10.5):
    for part in re.split(r"(\*\*.*?\*\*)", text):
        if not part:
            continue
        if part.startswith("**") and part.endswith("**"):
            font(paragraph.add_run(part[2:-2]), size, True)
        else:
            font(paragraph.add_run(part), size)


def page_number(paragraph):
    paragraph.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    font(paragraph.add_run("Page "), 9)
    run = paragraph.add_run()
    begin = OxmlElement("w:fldChar"); begin.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText"); instr.set(qn("xml:space"), "preserve"); instr.text = " PAGE "
    separate = OxmlElement("w:fldChar"); separate.set(qn("w:fldCharType"), "separate")
    result = OxmlElement("w:t"); result.text = "1"
    end = OxmlElement("w:fldChar"); end.set(qn("w:fldCharType"), "end")
    run._r.extend([begin, instr, separate, result, end])


def build():
    doc = Document()
    section = doc.sections[0]
    section.top_margin = Inches(0.65)
    section.bottom_margin = Inches(0.6)
    section.left_margin = Inches(0.72)
    section.right_margin = Inches(0.72)

    normal = doc.styles["Normal"]
    normal.font.name = "Aptos"; normal.font.size = Pt(10.5)
    normal.paragraph_format.space_after = Pt(5)
    for name, size in [("Title", 23), ("Heading 1", 15), ("Heading 2", 12)]:
        style = doc.styles[name]
        style.font.name = "Aptos Display"; style.font.size = Pt(size); style.font.bold = True
        style.font.color.rgb = RGBColor(0, 0, 0)
        style.paragraph_format.keep_with_next = True

    header = section.header.paragraphs[0]
    header.text = "Visa 801 outstanding items    Client reference WaKr-2631-801-3"
    header.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    for run in header.runs:
        font(run, 8.5); run.font.color.rgb = RGBColor(90, 90, 90)
    page_number(section.footer.paragraphs[0])

    for raw in SOURCE.read_text(encoding="utf-8").splitlines():
        if not raw.strip():
            continue
        if raw.startswith("# "):
            inline(doc.add_paragraph(style="Title"), raw[2:], 23)
        elif raw.startswith("## "):
            inline(doc.add_paragraph(style="Heading 1"), raw[3:], 15)
        elif raw.startswith("* ["):
            match = re.match(r"\* \[([ xX])\] (.*)", raw)
            checked = match.group(1).lower() == "x"
            p = doc.add_paragraph()
            p.paragraph_format.left_indent = Inches(0.22)
            p.paragraph_format.first_line_indent = Inches(-0.18)
            status = p.add_run("DONE  " if checked else "TO DO  ")
            font(status, 9.5, True)
            status.font.color.rgb = RGBColor(26, 112, 61) if checked else RGBColor(156, 72, 0)
            if not checked and "**" in match.group(2):
                status.font.highlight_color = WD_COLOR_INDEX.YELLOW
            inline(p, match.group(2), 10.2)
        elif raw.startswith("* "):
            p = doc.add_paragraph()
            p.paragraph_format.left_indent = Inches(0.22)
            p.paragraph_format.first_line_indent = Inches(-0.16)
            font(p.add_run("•  "), 10.5)
            inline(p, raw[2:], 10.2)
        else:
            inline(doc.add_paragraph(), raw.replace("  ", ""), 10.2)

    doc.core_properties.title = "Visa 801 Outstanding Items Report"
    doc.core_properties.subject = "Outstanding documents, questions and actions as at 18 October 2026"
    doc.core_properties.keywords = "Visa 801, outstanding items, WaKr-2631-801-3"
    doc.save(OUTPUT)
    print(OUTPUT)


if __name__ == "__main__":
    build()
